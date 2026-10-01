"""
report_excel.py

Builds the downloadable route report (.xlsx) from a /solve response:
stop level, trip level, truck level and an overall summary sheet.
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="12151A")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12)
KM_FORMAT = "0.00"
MONEY_FORMAT = '"$"#,##0.00'
PERCENT_FORMAT = "0.0%"
SIZE_COLUMNS = (("small_boxes", "Small"), ("mid_boxes", "Mid"), ("large_boxes", "Large"))

# Every trip currently starts and ends at one warehouse, so each truck runs
# exactly one trip; the trip/truck sheets are kept separate for when a truck
# can run several trips.
TRIP_NUMBER = 1


def _write_table(ws, start_row, headers, rows, formats=None):
    """Writes a styled header + rows; formats maps 0-based column -> number format."""
    formats = formats or {}
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start_row + 1):
        for col, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=col, value=value)
            if col - 1 in formats and isinstance(value, (int, float)):
                cell.number_format = formats[col - 1]
    return start_row + len(rows)


def _finish_sheet(ws, header_row=1, filter_table=True):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), len(str(cell.value)))
    for column, width in widths.items():
        ws.column_dimensions[column].width = min(max(10, width + 2), 60)
    if filter_table:
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        ws.auto_filter.ref = ws.dimensions


def build_route_report(result, country=None, dataset=None, settings=None):
    settings = settings or {}
    capacity = settings.get("vehicle_capacity")

    stop_rows, trip_rows = [], []
    trucks = {}
    totals_by_size = {key: 0 for key, _ in SIZE_COLUMNS}

    for route in result["truck_routes"]:
        details = route["stop_details"]
        warehouse = details[0]
        truck_id = route["truck"]
        route_country = next((d["country"] for d in details if d.get("country")), None) or country
        trip_sizes = {key: 0 for key, _ in SIZE_COLUMNS}
        stations = []

        for d in details:
            if d["is_depot"]:
                continue
            stations.append(d["name"])
            sizes = [(key, label, d.get(key) or 0) for key, label in SIZE_COLUMNS]
            if not any(count for _, _, count in sizes):
                sizes = [(None, "Unspecified", d["batteries"])]
            for key, label, count in sizes:
                if count <= 0:
                    continue
                if key:
                    trip_sizes[key] += count
                stop_rows.append([
                    truck_id, TRIP_NUMBER, d.get("country") or route_country,
                    warehouse["code"], warehouse["name"], d["code"], d["name"],
                    label, count, None, None,
                    d["cumulative_km"], d["cumulative_fuel_cost"],
                    len(stations), d["leg_km"], d["leg_fuel_cost"], d["cumulative_hours"],
                    d.get("region"), d.get("district"),
                    d.get("station_faulty_batteries"), d.get("working_batteries"),
                ])

        batteries = route["total_batteries_picked"]
        fill = batteries / capacity if capacity else None
        trip_rows.append([
            truck_id, TRIP_NUMBER, route_country, warehouse["code"], warehouse["name"],
            len(stations), trip_sizes["small_boxes"], trip_sizes["mid_boxes"], trip_sizes["large_boxes"],
            batteries, capacity, fill, route["distance_km"], route["driving_hours"],
            route["estimated_days"], route["fuel_cost"],
            " > ".join([warehouse["name"], *stations, warehouse["name"]]),
        ])
        for key in totals_by_size:
            totals_by_size[key] += trip_sizes[key]

        truck = trucks.setdefault(truck_id, {
            "trips": 0, "warehouses": [], "stations": 0, "batteries": 0,
            "sizes": {key: 0 for key, _ in SIZE_COLUMNS}, "km": 0.0, "hours": 0.0, "fuel": 0.0,
        })
        truck["trips"] += 1
        if warehouse["name"] not in truck["warehouses"]:
            truck["warehouses"].append(warehouse["name"])
        truck["stations"] += len(stations)
        truck["batteries"] += batteries
        for key in truck["sizes"]:
            truck["sizes"][key] += trip_sizes[key]
        truck["km"] += route["distance_km"]
        truck["hours"] += route["driving_hours"]
        truck["fuel"] += route["fuel_cost"]

    wb = Workbook()

    ws = wb.active
    ws.title = "stop level report"
    _write_table(ws, 1, [
        "Truck_ID", "Trip_Number", "Country", "warehouse_ID", "warehouse_Name", "Station_ID", "Station_Name",
        "Battery Size type", "faulty Batteries_Delivered", "Battery Manf Name", "Battery Dimension (if available)",
        "Final distance in km", "Fuel Cost in $",
        "Stop_Sequence", "Leg distance in km", "Leg fuel cost in $", "Cumulative driving hours",
        "Region", "District",
        "Station faulty batteries", "Working batteries", "Total batteries (faulty + working)", "Faulty battery %",
    ], stop_rows, {11: KM_FORMAT, 12: MONEY_FORMAT, 14: KM_FORMAT, 15: MONEY_FORMAT})
    # Total and % are formulas so they recalculate when working batteries are
    # filled in by hand; blank working count -> blank total and % (not 100%).
    for r in range(2, len(stop_rows) + 2):
        ws[f"V{r}"] = f'=IF(U{r}="","",T{r}+U{r})'
        ws[f"W{r}"] = f'=IF(OR(U{r}="",V{r}=0),"",T{r}/V{r})'
        ws[f"W{r}"].number_format = PERCENT_FORMAT
    _finish_sheet(ws)

    ws = wb.create_sheet("truck level")
    _write_table(ws, 1, [
        "Truck_ID", "Trips", "Warehouse(s)", "Stations visited", "Small", "Mid", "Large",
        "Total batteries", "Truck capacity per trip", "Average fill %", "Total distance in km",
        "Driving hours", "Fuel Cost in $",
    ], [
        [
            truck_id, t["trips"], ", ".join(t["warehouses"]), t["stations"],
            t["sizes"]["small_boxes"], t["sizes"]["mid_boxes"], t["sizes"]["large_boxes"],
            t["batteries"], capacity,
            t["batteries"] / (capacity * t["trips"]) if capacity else None,
            round(t["km"], 2), round(t["hours"], 2), round(t["fuel"], 2),
        ]
        for truck_id, t in trucks.items()
    ], {9: PERCENT_FORMAT, 10: KM_FORMAT, 12: MONEY_FORMAT})
    _finish_sheet(ws)

    ws = wb.create_sheet("trip level")
    _write_table(ws, 1, [
        "Truck_ID", "Trip_Number", "Country", "warehouse_ID", "warehouse_Name", "Stations visited",
        "Small", "Mid", "Large", "Total batteries", "Truck capacity", "Fill %",
        "Total distance in km", "Driving hours", "Estimated days", "Fuel Cost in $", "Route",
    ], trip_rows, {11: PERCENT_FORMAT, 12: KM_FORMAT, 15: MONEY_FORMAT})
    _finish_sheet(ws)

    ws = wb.create_sheet("Overall Report")
    total_batteries = result["total_cartons_picked"]
    summary = [
        ("Generated at", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Country", country or ""),
        ("Dataset", dataset or ""),
        ("Routing engine", result["routing_algorithm"]),
        ("Truck capacity (batteries)", capacity),
        ("Fill target", settings.get("target_fill_ratio")),
        ("Max driving hours per day", settings.get("max_driving_hours")),
        ("Max trip days", settings.get("max_trip_days")),
        ("Fuel cost per km ($)", settings.get("road_cost_per_km")),
        ("Excluded regions", ", ".join(settings.get("excluded_regions") or []) or "None"),
        ("Excluded districts", ", ".join(settings.get("excluded_districts") or []) or "None"),
        ("Stops removed by excluded areas", len(result.get("region_excluded_stops", []))),
        ("Trucks used", result["trucks_needed"]),
        ("Trips", len(result["truck_routes"])),
        ("Stations served", result["num_stops_served"]),
        ("Stations dropped", len(result["dropped_stops"])),
        ("Small batteries collected", totals_by_size["small_boxes"]),
        ("Mid batteries collected", totals_by_size["mid_boxes"]),
        ("Large batteries collected", totals_by_size["large_boxes"]),
        ("Total batteries collected", total_batteries),
        ("Overall fill %", result["fill_ratio"]),
        ("Fill target met", "Yes" if result["meets_fill_target"] else "No"),
        ("Total distance in km", round(sum(r["distance_km"] for r in result["truck_routes"]), 2)),
        ("Total driving hours", round(sum(r["driving_hours"] for r in result["truck_routes"]), 2)),
        ("Fuel cost ($)", result["fuel_cost"]),
        ("Toll tax ($)", result["toll_tax_total"]),
        ("Total operating cost ($)", result["total_operating_cost"]),
        ("Cost per battery ($)", round(result["total_operating_cost"] / total_batteries, 2) if total_batteries else None),
    ]
    row = _write_table(ws, 1, ["Metric", "Value"], [list(item) for item in summary])
    for r in range(2, row + 1):
        label = ws.cell(row=r, column=1).value
        value_cell = ws.cell(row=r, column=2)
        if label in ("Fill target", "Overall fill %") and isinstance(value_cell.value, (int, float)):
            value_cell.number_format = PERCENT_FORMAT
        elif "($)" in label and isinstance(value_cell.value, (int, float)):
            value_cell.number_format = MONEY_FORMAT

    dropped = result.get("dropped_stop_details", [])
    ws.cell(row=row + 2, column=1, value="Dropped stations (not served by any truck)").font = SECTION_FONT
    if dropped:
        row = _write_table(ws, row + 3, ["Station_ID", "Station_Name", "Country", "Region", "District", "Small", "Mid", "Large", "Total batteries"], [
            [d["code"], d["name"], d.get("country") or country, d.get("region"), d.get("district"),
             d["small_boxes"], d["mid_boxes"], d["large_boxes"], d["batteries"]]
            for d in dropped
        ])
    else:
        row += 3
        ws.cell(row=row, column=1, value="None - every station was served.")

    area_excluded = result.get("region_excluded_stops", [])
    ws.cell(row=row + 2, column=1, value="Stops removed by excluded regions/districts").font = SECTION_FONT
    if area_excluded:
        _write_table(ws, row + 3, ["Code", "Name", "Type", "Region", "District", "Total batteries"], [
            [d["code"], d["name"], "Warehouse" if d["is_depot"] else "Station", d.get("region"), d.get("district"), d["batteries"]]
            for d in area_excluded
        ])
    else:
        ws.cell(row=row + 3, column=1, value="None - no areas were excluded.")
    _finish_sheet(ws, filter_table=False)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
